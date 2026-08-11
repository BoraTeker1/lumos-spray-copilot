"""Load a downloaded USDA PDP annual release into the residue reference table.

    python -m app.pdp_sync /path/to/2016PDPDatabase.zip --crops strawberry,tomato

Releases: https://www.ams.usda.gov/datasets/pdp/pdpdata  (1992-2024, one zip per year)

EXPLICIT on purpose, exactly like `app/label_sync.py`, and for the same two reasons:
seeding this would let a demo farm show national residue statistics beside simulated
sprays as though the two were the same kind of fact, and loading it at startup would mean
a reference dataset changed with no record of who loaded what.

Takes a LOCAL PATH rather than a URL. The loader does not download: it reads bytes a
human fetched, so the release under a digest is the release someone actually looked at,
and this module needs no network access, no retry policy, and no place to put an API key.

The commodity filter defaults to the crops this product covers. Loading all ~20 sampled
commodities is supported (`--crops all`) but is not the default, because a residue profile
for a crop the engine has no other data about invites reading it as coverage.
"""
from __future__ import annotations

import argparse
import io
import re
import sys
import zipfile
from typing import Any

from app import crop_aliases, crud, pdp_dataset
from app.database import SessionLocal

# PDP publishes concentration/tolerance units as single letters.
TOLERANCE_UNITS = {"M": "ppm", "B": "ppb", "T": "ppt"}
# Non-numeric tolerance bases. Preserved as codes; NEVER coerced to a number and never
# treated as "no limit" — see models.ResidueReferenceRecord.
TOLERANCE_BASES = {"NT", "EX", "SU"}


class PdpReleaseError(ValueError):
    """The zip did not look like a PDP annual release."""


def _member(zf: zipfile.ZipFile, pattern: str) -> str:
    hits = [n for n in zf.namelist() if re.search(pattern, n, re.I)]
    if not hits:
        raise PdpReleaseError(
            f"No file matching {pattern!r} in the archive. Expected a USDA PDP annual "
            f"release zip; found: {sorted(zf.namelist())}"
        )
    return hits[0]


def _program_year(zf: zipfile.ZipFile) -> int:
    name = _member(zf, r"Samples\.txt$")
    m = re.search(r"PDP(\d{2})Samples", name, re.I)
    if not m:
        raise PdpReleaseError(f"Cannot read the program year from {name!r}.")
    yy = int(m.group(1))
    # PDP began in 1991; two-digit years below that belong to the 2000s.
    return 1900 + yy if yy >= 91 else 2000 + yy


def _sheet_rows(zf: zipfile.ZipFile) -> dict[str, list[list[Any]]]:
    """Read the release's reference workbook into {sheet_name: rows}.

    Handles both forms USDA has shipped: .xlsx (2020+) and legacy .xls (2019 and
    earlier). Both readers are imported lazily so neither is needed to run the API — the
    same technique `ingest/cimis.py` uses for httpx, and for the same reason: an optional
    capability must not make the module unimportable where it is unused.
    """
    try:
        name = _member(zf, r"ReferenceTables.*\.xlsx$")
    except PdpReleaseError:
        name = _member(zf, r"ReferenceTables.*\.xls$")

    payload = zf.read(name)
    out: dict[str, list[list[Any]]] = {}

    if name.lower().endswith(".xlsx"):
        try:
            import openpyxl
        except ImportError as exc:  # pragma: no cover - environment-dependent
            raise PdpReleaseError(
                "Reading a modern PDP release needs openpyxl (`pip install openpyxl`)."
            ) from exc
        wb = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        for sheet in wb.sheetnames:
            out[sheet.strip().lower()] = [
                list(r) for r in wb[sheet].iter_rows(values_only=True)
            ]
    else:
        try:
            import xlrd
        except ImportError as exc:  # pragma: no cover - environment-dependent
            raise PdpReleaseError(
                "Reading a pre-2020 PDP release needs xlrd (`pip install xlrd`)."
            ) from exc
        bk = xlrd.open_workbook(file_contents=payload)
        for sheet in bk.sheet_names():
            sh = bk.sheet_by_name(sheet)
            out[sheet.strip().lower()] = [
                [sh.cell_value(r, c) for c in range(sh.ncols)] for r in range(sh.nrows)
            ]
    return out


def _cell(value: Any) -> str:
    """Workbook cell to a trimmed string, without turning 11.0 into '11.0'.

    xlrd returns every numeric cell as a float, so the pesticide code 011 arrives as 11.0.
    Codes are zero-padded text in the data files; losing the padding would silently fail
    every join against Results.txt.
    """
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _lookup_map(rows: list[list[Any]], code_len: tuple[int, ...]) -> dict[str, str]:
    """Extract {code: name} from a reference sheet by ROW SHAPE, not by header text.

    Header wording drifts between releases ("Extract" vs "Extraction", title rows that
    repeat the word "Pesticide"), so matching on it produced a reader that worked on 2024
    and silently returned nothing on 2016. Shape is stable: a data row is a short code
    followed by a non-numeric name.
    """
    out: dict[str, str] = {}
    for row in rows:
        if len(row) < 2:
            continue
        code, name = _cell(row[0]), _cell(row[1])
        if not code or not name:
            continue
        if len(code) not in code_len or " " in code:
            continue
        if re.fullmatch(r"[\d.]+", name):
            continue
        out.setdefault(code.upper(), name)
    return out


def _pad(code: str, width: int) -> str:
    """Zero-pad a numeric code back to its fixed width (11 -> '011')."""
    return code.zfill(width) if code.isdigit() else code


def read_reference(zf: zipfile.ZipFile) -> tuple[dict[str, str], dict[str, str], dict]:
    """Commodity names, pesticide names, and EPA tolerances from the release workbook."""
    sheets = _sheet_rows(zf)

    commodity_rows = sheets.get("commodity", [])
    pest_rows = sheets.get("pest code", []) or sheets.get("pestcode", [])
    tol_rows = sheets.get("tolerance", [])

    commodities = _lookup_map(commodity_rows, code_len=(2,))
    pesticides = {
        _pad(k, 3): v for k, v in _lookup_map(pest_rows, code_len=(2, 3)).items()
    }

    tolerances: dict[tuple[str, str], tuple[float | None, str | None, str | None]] = {}
    for row in tol_rows:
        if len(row) < 3:
            continue
        pest, commod, raw = _pad(_cell(row[0]), 3), _cell(row[1]).upper(), _cell(row[2])
        unit_code = _cell(row[3]).upper() if len(row) > 3 else ""
        if not pest or len(commod) != 2 or not raw:
            continue
        value: float | None = None
        basis: str | None = None
        if raw.upper() in TOLERANCE_BASES:
            basis = raw.upper()
        else:
            try:
                value = float(raw)
            except ValueError:
                # An unrecognised tolerance string is dropped rather than guessed at.
                continue
        tolerances[(pest, commod)] = (
            value, basis, TOLERANCE_UNITS.get(unit_code) if value is not None else None,
        )

    if not commodities or not pesticides:
        raise PdpReleaseError(
            "The release workbook yielded no commodity or pesticide names. Refusing to "
            "load aggregates that could not be labelled — codes alone are unmatchable."
        )
    return commodities, pesticides, tolerances


def _commodity_codes_for(crops: list[str], commodities: dict[str, str]) -> set[str]:
    """Resolve crop names to PDP commodity codes through the curated alias table.

    Exact or curated alias only. A crop that matches nothing is reported rather than
    quietly dropped, because "I loaded strawberries" silently loading nothing is the
    failure this whole codebase is built to avoid.
    """
    wanted: set[str] = set()
    unmatched: list[str] = []
    for crop in crops:
        hits = {
            code for code, name in commodities.items()
            if crop_aliases.match_crops(crop, name) == crop_aliases.MATCH
        }
        if hits:
            wanted |= hits
        else:
            unmatched.append(crop)
    if unmatched:
        print(
            f"note: no commodity in this release matches {unmatched}. PDP samples about "
            f"20 commodities a year; this release covers: "
            f"{sorted(commodities.values())}"
        )
    return wanted


def load(
    zip_path: str,
    *,
    crops: list[str] | None,
    domestic_only: bool = True,
    loaded_by: str | None = None,
) -> dict:
    with zipfile.ZipFile(zip_path) as zf:
        year = _program_year(zf)
        commodities, pesticides, tolerances = read_reference(zf)

        if crops is None:
            codes = None
            scope = "all sampled commodities"
        else:
            codes = _commodity_codes_for(crops, commodities)
            scope = ", ".join(sorted(commodities[c] for c in codes)) or "(none)"
            if not codes:
                return {
                    "program_year": year, "scope": scope, "aggregates_read": 0,
                    "records_created": 0, "unchanged": 0, "skipped_unnamed": 0,
                }

        samples_name = _member(zf, r"Samples\.txt$")
        results_name = _member(zf, r"Results\.txt$")
        samples_bytes = zf.read(samples_name)
        results_bytes = zf.read(results_name)

        digest = pdp_dataset.source_digest(samples_bytes, results_bytes)
        samples = pdp_dataset.parse_samples(samples_bytes.decode("latin-1").splitlines())
        aggregates = pdp_dataset.aggregate(
            samples,
            pdp_dataset.iter_results(results_bytes.decode("latin-1").splitlines()),
            program_year=year,
            commodity_codes=codes,
            domestic_only=domestic_only,
        )

    source_reference = (
        f"USDA AMS Pesticide Data Program {year} Annual Database — "
        f"https://www.ams.usda.gov/datasets/pdp/pdpdata "
        f"({samples_name}, {results_name}; "
        f"{'domestic (ORIGIN=1) samples only' if domestic_only else 'all sample origins'})"
    )

    db = SessionLocal()
    try:
        result = crud.load_residue_reference(
            db,
            aggregates,
            commodity_names=commodities,
            pesticide_names=pesticides,
            tolerances=tolerances,
            source_reference=source_reference,
            source_digest=digest,
            loaded_by=loaded_by,
        )
    finally:
        db.close()

    result.update({"program_year": year, "scope": scope, "source_digest": digest})
    return result


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("zip_path", help="Path to a downloaded <year>PDPDatabase.zip")
    parser.add_argument(
        "--crops", default="strawberry,tomato",
        help="Comma-separated crop names, or 'all'. Default: the wedge crops.",
    )
    parser.add_argument(
        "--all-origins", action="store_true",
        help="Include imported samples. Default is US-grown (ORIGIN=1) only, which is "
             "what a California grower's question is about.",
    )
    parser.add_argument("--loaded-by", default=None, help="Attribution for the load.")
    args = parser.parse_args(argv)

    crops = None if args.crops.strip().lower() == "all" else [
        c.strip() for c in args.crops.split(",") if c.strip()
    ]

    try:
        result = load(
            args.zip_path, crops=crops, domestic_only=not args.all_origins,
            loaded_by=args.loaded_by,
        )
    except (PdpReleaseError, pdp_dataset.PdpFormatError) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 1

    print(
        f"program year:     {result['program_year']}\n"
        f"commodities:      {result['scope']}\n"
        f"aggregates read:  {result['aggregates_read']}\n"
        f"records created:  {result['records_created']}\n"
        f"already loaded:   {result['unchanged']}\n"
        f"skipped (unnamed):{result['skipped_unnamed']}"
    )
    if result["records_created"] == 0 and result["unchanged"] == 0:
        print(
            "\nNothing was loaded. Every residue lookup will keep refusing with "
            "`no_residue_reference_loaded` or `crop_not_in_program`, which is the "
            "correct state until a release covering your crop is loaded. PDP rotates "
            "commodities: fresh strawberries were last sampled in 2016."
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
